"""Convierte Files/productos.xlsx en src/data/mock/<marca>.products.json (uso puntual).

Uso: python scripts/excel-to-json.py [ruta_excel] [marca]
"""
import json, re, sys, unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
excel = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / "Files" / "productos.xlsx"
brand = sys.argv[2] if len(sys.argv) > 2 else "audi"

CATEGORIES = [  # (categoría, palabras clave en minúsculas)
    ("Relojes", ["reloj", "cron"]),
    ("Gafas", ["gafas"]),
    ("Gorras", ["gorra"]),
    ("Paraguas", ["paraguas", "umbrella"]),
    ("Llaveros", ["llavero"]),
    ("Ropa", ["chaqueta", "polo", "camiseta", "chompa"]),
    ("Bolsos y viaje", ["bolso", "bolsa", "backpack", "bag", "maleta", "neceser", "canguro", "billetera", "airpods"]),
    ("Escritura", ["esferografico", "esferográfico", "biros"]),
    ("Bebidas y hogar", ["botella", "thermo", "mug", "taza", "tomatoddo", "caja frigorifica", "caja frigorífica"]),
    ("Coleccionables", ["modelo escala", "carro audi"]),
]

def slugify(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")

def categorize(name: str) -> str:
    n = name.lower()
    for cat, keys in CATEGORIES:
        if any(k in n for k in keys):
            return cat
    return "Accesorios"

ws = load_workbook(excel).active
products, seen = [], {}
for name, price, desc, image in list(ws.iter_rows(min_row=2, values_only=True)):
    if not name:
        continue
    slug = slugify(name)
    seen[slug] = seen.get(slug, 0) + 1
    if seen[slug] > 1:
        slug = f"{slug}-{seen[slug]}"
    products.append({
        "id": f"{brand}-{image.split('.')[0]}",
        "slug": slug,
        "brand": brand,
        "name": name,
        "description": desc or "",
        "price": {"amount": float(price), "currency": "USD"},
        "image": {"src": f"/brands/{brand}/products/{image}", "alt": name},
        "category": categorize(name),
    })

out = ROOT / "src" / "data" / "mock" / f"{brand}.products.json"
out.write_text(json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"{len(products)} productos -> {out}")
from collections import Counter
print(Counter(p["category"] for p in products))
